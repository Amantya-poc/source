"""Generate ASP Console architecture Word document with diagram."""

from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOCS_DIR = Path(__file__).resolve().parent
DIAGRAM_PATH = DOCS_DIR / "architecture-diagram.png"
OUTPUT_PATH = DOCS_DIR / "ASP-Console-Architecture-Documentation.docx"


def get_font(size: int, bold: bool = False):
    try:
        name = "arialbd.ttf" if bold else "arial.ttf"
        return ImageFont.truetype(name, size)
    except OSError:
        return ImageFont.load_default()


def draw_architecture_diagram(path: Path) -> None:
    width, height = 1400, 900
    img = Image.new("RGB", (width, height), "#f8fafc")
    draw = ImageDraw.Draw(img)

    title_font = get_font(28, bold=True)
    box_font = get_font(18, bold=True)
    small_font = get_font(14)
    tiny_font = get_font(12)

    draw.text((width // 2, 30), "ASP Console — System Architecture", fill="#1e3a6b", anchor="mt", font=title_font)

    def box(x, y, w, h, fill, title, lines, border="#334155"):
        draw.rounded_rectangle([x, y, x + w, y + h], radius=12, fill=fill, outline=border, width=2)
        draw.text((x + w // 2, y + 18), title, fill="#0f172a", anchor="mt", font=box_font)
        ty = y + 52
        for line in lines:
            draw.text((x + w // 2, ty), line, fill="#334155", anchor="mt", font=small_font)
            ty += 22

    # Browser layer
    box(80, 90, 1240, 170, "#dbeafe", "Angular Frontend (localhost:4200)", [
        "DashboardComponent — sidebar, controls, timeline",
        "AerialDeviceMapComponent — OpenLayers map, simulation",
        "PlanPanelComponent | OpDashPanelComponent | Deploy areas",
    ])

    # Services row
    box(120, 300, 360, 130, "#fef3c7", "GeoServer :8080", [
        "WMS map tiles",
        "ne:NE2_HR_LC_SR_W_DR",
    ])
    box(520, 300, 360, 130, "#dcfce7", "Spring Boot API :8081", [
        "PlanController | DeployAreaController",
        "PlanService | TimelineService",
    ])
    box(920, 300, 360, 130, "#ede9fe", "PostgreSQL :5433", [
        "Database: drdo_poc",
        "plans | timeline_settings | deploy_areas",
    ])

    # Arrows helper
    def arrow(x1, y1, x2, y2, color="#64748b"):
        draw.line([x1, y1, x2, y2], fill=color, width=3)
        if y2 > y1:
            draw.polygon([(x2, y2), (x2 - 8, y2 - 14), (x2 + 8, y2 - 14)], fill=color)
        elif x2 > x1:
            draw.polygon([(x2, y2), (x2 - 14, y2 - 8), (x2 - 14, y2 + 8)], fill=color)
        elif x2 < x1:
            draw.polygon([(x2, y2), (x2 + 14, y2 - 8), (x2 + 14, y2 + 8)], fill=color)

    arrow(700, 260, 300, 300, "#2563eb")
    draw.text((420, 268), "WMS tiles", fill="#2563eb", font=tiny_font)
    arrow(700, 260, 700, 300, "#16a34a")
    draw.text((710, 268), "REST JSON", fill="#16a34a", font=tiny_font)
    arrow(880, 365, 920, 365, "#7c3aed")
    draw.text((885, 345), "JPA", fill="#7c3aed", font=tiny_font)

    # Flow boxes
    box(80, 480, 380, 150, "#ffffff", "1. Planning Flow", [
        "User draws route on map → POST /api/plans",
        "Backend stores plan, recalculates timeline",
    ])
    box(510, 480, 380, 150, "#ffffff", "2. Simulation Flow", [
        "Play → local animation loop (no API)",
        "Aircraft move by distance, speed, start time",
    ])
    box(940, 480, 380, 150, "#ffffff", "3. Timeline Slider", [
        "Start = min(plan start dates)",
        "End = max(plan end dates); scrub = local",
    ])

    box(295, 680, 810, 150, "#ffffff", "Fuel & Duration Rules (per plan)", [
        "Capacity: 1000 L  |  Consumption: 1 L/min  |  Max plans: 3",
        "Travel duration = route distance ÷ speed (capped at 1000 min)",
        "Plan end = startingDate + travelDuration",
    ])

    img.save(path, "PNG")


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_bullet(doc: Document, text: str, bold_prefix: str = "") -> None:
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
        p.add_run(text)
    else:
        p.add_run(text)


def build_document(diagram_path: Path, output_path: Path) -> None:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("ASP Console — Architecture & Flow Documentation", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run("Aerial Simulation Planning (ASP) Console\n")
    run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x6B)
    meta.add_run("Project: issapoc/source | Stack: Angular 21 + Spring Boot + PostgreSQL + GeoServer + OpenLayers")

    doc.add_paragraph()

    add_heading(doc, "1. Executive Summary", 1)
    doc.add_paragraph(
        "ASP Console is a full-stack geospatial application for planning up to three aircraft "
        "mission routes, simulating movement over time on an interactive map, managing deploy "
        "areas, and analyzing fuel consumption. The frontend is an Angular single-page dashboard; "
        "the backend is a Spring Boot REST API backed by PostgreSQL. Map imagery is served by GeoServer via WMS."
    )

    add_heading(doc, "2. System Architecture Diagram", 1)
    doc.add_picture(str(diagram_path), width=Inches(6.5))
    last_paragraph = doc.paragraphs[-1]
    last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_heading(doc, "3. Technology Stack", 1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Layer"
    hdr[1].text = "Technology"
    hdr[2].text = "Purpose"
    rows = [
        ("Frontend", "Angular 21, PrimeNG 21", "UI, forms, charts, routing"),
        ("Map", "OpenLayers 10", "2D map, routes, aircraft markers, deploy ellipses"),
        ("Backend", "Spring Boot, JPA", "REST API, business logic, persistence"),
        ("Database", "PostgreSQL (drdo_poc)", "Plans, timeline settings, deploy areas"),
        ("Map server", "GeoServer :8080", "WMS base map tiles"),
        ("API port", "Spring Boot :8081", "Plans and deploy-area endpoints"),
    ]
    for layer, tech, purpose in rows:
        row = table.add_row().cells
        row[0].text = layer
        row[1].text = tech
        row[2].text = purpose

    add_heading(doc, "4. Application Structure", 1)
    add_bullet(doc, " — root route loads ShellComponent → DashboardComponent", "Routing:")
    add_bullet(doc, " — collapsible sidebar with Planning, Deployment, Analysis, History tabs", "Dashboard layout:")
    add_bullet(doc, " — OpenLayers map with top simulation controls and bottom timeline slider", "Map area:")
    add_bullet(doc, " — aircraft status cards (speed, fuel, route progress)", "Sidebar footer:")

    add_heading(doc, "5. Backend API Endpoints", 1)
    api_table = doc.add_table(rows=1, cols=3)
    api_table.style = "Table Grid"
    api_hdr = api_table.rows[0].cells
    api_hdr[0].text = "Method"
    api_hdr[1].text = "Endpoint"
    api_hdr[2].text = "Description"
    apis = [
        ("GET", "/api/plans/dashboard", "Load plans, route events, timeline in one call"),
        ("GET", "/api/plans/timeline", "Slider start/end times"),
        ("POST", "/api/plans", "Save new plan (max 3)"),
        ("DELETE", "/api/plans", "Clear all plans"),
        ("GET", "/api/deploy-areas", "List deploy ellipses"),
        ("POST", "/api/deploy-areas", "Create deploy area"),
        ("PUT", "/api/deploy-areas/{key}", "Update deploy area geometry"),
        ("DELETE", "/api/deploy-areas/{key}", "Delete deploy area"),
    ]
    for method, endpoint, desc in apis:
        row = api_table.add_row().cells
        row[0].text = method
        row[1].text = endpoint
        row[2].text = desc

    add_heading(doc, "6. Core Data Flows", 1)

    add_heading(doc, "6.1 Startup / Page Load", 2)
    doc.add_paragraph(
        "On dashboard init, GET /api/plans/dashboard returns all saved plans, computed route "
        "intersection events, and timeline bounds. The frontend restores routes on the map, "
        "configures simulation vehicles, and sets the timeline slider range."
    )

    add_heading(doc, "6.2 Planning Flow", 2)
    steps = [
        "User enters plan name, speed, and start date/time in the Planning tab.",
        "User clicks Map Route and adds waypoints by clicking on the map.",
        "User clicks Finish → POST /api/plans with route waypoints.",
        "Backend computes route distance (haversine), stores plan in PostgreSQL, recalculates timeline.",
        "Frontend adds the plan, starts a simulation vehicle, and reloads timeline settings.",
    ]
    for i, step in enumerate(steps, 1):
        add_bullet(doc, step)

    add_heading(doc, "6.3 Simulation Flow", 2)
    doc.add_paragraph(
        "When the user clicks Play, a requestAnimationFrame loop runs entirely in the browser. "
        "Each aircraft waits until its own startingDate offset, then moves along its route at "
        "the configured speed. Fuel decreases based on distance traveled. No API calls occur during playback."
    )

    add_heading(doc, "6.4 Timeline Slider", 2)
    doc.add_paragraph("Timeline bounds are computed from all saved plans:")
    add_bullet(doc, " — earliest startingDate among all plans", "Slider START")
    add_bullet(doc, " — latest plan end time among all plans", "Slider END")
    add_bullet(doc, " — startingDate + travelDuration (distance ÷ speed, capped at 1000 min)", "Each plan end")
    doc.add_paragraph(
        "Moving the slider forward or backward recalculates aircraft position and fuel locally in "
        "the browser. It does NOT call the backend API. The API is only used on page load, after "
        "saving a plan, or after clearing plans."
    )

    add_heading(doc, "7. Fuel & Duration Rules", 1)
    fuel_table = doc.add_table(rows=1, cols=2)
    fuel_table.style = "Table Grid"
    fuel_table.rows[0].cells[0].text = "Rule"
    fuel_table.rows[0].cells[1].text = "Value"
    fuel_rules = [
        ("Fuel capacity (per plan)", "1000 liters"),
        ("Fuel consumption", "1 liter per minute of travel time"),
        ("Max flight time (fuel cap)", "1000 minutes"),
        ("Travel duration", "route distance ÷ speed (capped by fuel limit)"),
        ("Fuel during simulation", "1000 L − (distance traveled converted to minutes × 1 L/min)"),
        ("Maximum plans", "3"),
    ]
    for rule, value in fuel_rules:
        row = fuel_table.add_row().cells
        row[0].text = rule
        row[1].text = value

    add_heading(doc, "8. Route Events & Analysis", 1)
    doc.add_paragraph(
        "Route intersection events are computed when two plan route segments cross. These appear "
        "as markers on the map and as snap points on the timeline in Event playback mode. The "
        "Analysis tab (O/P dash) shows Chart.js fuel charts derived from plan routes and events."
    )

    add_heading(doc, "9. Deployment Areas", 1)
    doc.add_paragraph(
        "In the Deployment tab, users place elliptical deploy areas on the map. Areas are "
        "persisted via /api/deploy-areas and can be moved, resized, or reshaped. Geometry is "
        "stored in PostgreSQL and rendered on a separate OpenLayers vector layer."
    )

    add_heading(doc, "10. Key Source Files", 1)
    files = [
        ("frontend/src/app/pages/dashboard/dashboard.component.ts", "Main dashboard orchestration"),
        ("frontend/src/app/pages/dashboard/aerial-device-map/aerial-device-map.component.ts", "Map, simulation, deploy areas"),
        ("frontend/src/app/services/plan-api.service.ts", "REST client for plans"),
        ("backend/.../controller/PlanController.java", "Plan REST endpoints"),
        ("backend/.../service/TimelineService.java", "Timeline start/end calculation"),
        ("backend/.../util/PlanTimelineUtil.java", "Travel duration and plan end time"),
        ("backend/.../util/RouteDistanceUtil.java", "Route length (haversine)"),
        ("frontend/src/app/utils/plan-fuel.util.ts", "Fuel calculations"),
    ]
    for path, desc in files:
        add_bullet(doc, f" — {desc}", path)

    add_heading(doc, "11. Running the Application", 1)
    add_bullet(doc, "Frontend: cd frontend && npm start → http://localhost:4200")
    add_bullet(doc, "Backend: cd backend && mvn spring-boot:run → http://localhost:8081")
    add_bullet(doc, "Swagger UI: http://localhost:8081/swagger-ui.html")
    add_bullet(doc, "GeoServer: http://localhost:8080/geoserver (configure in geoserver.config.ts)")
    add_bullet(doc, "PostgreSQL: localhost:5433, database drdo_poc")

    doc.add_paragraph()
    footer = doc.add_paragraph("Document generated from issapoc/source project documentation.")
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.italic = True
    footer.runs[0].font.size = Pt(9)

    doc.save(output_path)


def main() -> None:
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    draw_architecture_diagram(DIAGRAM_PATH)
    build_document(DIAGRAM_PATH, OUTPUT_PATH)
    print(f"Created: {OUTPUT_PATH}")
    print(f"Diagram: {DIAGRAM_PATH}")


if __name__ == "__main__":
    main()
